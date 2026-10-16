from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


@dataclass
class ParsedBlock:
    page: int
    text: str


@dataclass
class ParsedDocument:
    method: str
    blocks: list[ParsedBlock]
    error: str = ""


def clean_text(text: str) -> str:
    text = (text or "").replace("\u00a0", " ").replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_text_file(path: Path) -> ParsedDocument:
    content = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            content = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise ValueError("无法识别文本编码")
    return ParsedDocument("text", [ParsedBlock(1, clean_text(content))])


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document

    document = Document(path)
    lines = []
    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if text:
            lines.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"\n表格 {table_index}")
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            lines.append(" | ".join(cells))
    return ParsedDocument("docx", [ParsedBlock(1, "\n".join(lines))])


def parse_pdf(path: Path) -> ParsedDocument:
    import pymupdf

    blocks = []
    with pymupdf.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            text = clean_text(page.get_text("text"))
            if text:
                blocks.append(ParsedBlock(page_number, text))
    if not blocks:
        return ParsedDocument(
            "pdf_ocr_required",
            [],
            "未提取到文字层，需要安装并接入 OCR 处理扫描版 PDF。",
        )
    return ParsedDocument("pdf_text", blocks)


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"暂不支持文件类型：{suffix}")
    if suffix in {".txt", ".md"}:
        return parse_text_file(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"暂不支持文件类型：{suffix}")
